"""Builds a single multi-sheet .xlsx workbook covering every major record in the
hospital — patients, staff (split by role), billing (consult/pharmacy/labs),
salaries, stock, and configuration catalogues — for the admin's one-click export.
"""

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.roles import Role
from app.modules.admin.models import (
    AuditLog,
    Department,
    MedicineFormulary,
    StaffSalary,
    TestCatalogue,
)
from app.modules.auth.models import User
from app.modules.doctors.models import Doctor, DoctorRoster
from app.modules.labs.models import TestOrder
from app.modules.patients.models import Patient
from app.modules.pharmacy.models import BillingEntry, DispenseLog, StockItem

HEADER_FONT = Font(bold=True)
DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def _fmt(value: datetime | date | None) -> str | None:
    if value is None:
        return None
    return value.isoformat()


def _write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> Worksheet:
    ws = wb.create_sheet(title=title[:31])  # Excel sheet-name length limit
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(row)
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(str(header)) + 4))
    ws.freeze_panes = "A2"
    return ws


def build_full_export(db: Session) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)

    dept_name_by_id = {d.id: d.name for d in db.execute(select(Department)).scalars()}

    # ---- Patients -----------------------------------------------------------
    patients = list(db.execute(select(Patient).order_by(Patient.created_at.desc())).scalars())
    _write_sheet(
        wb,
        "Patients",
        [
            "ID", "Full Name", "DOB", "Gender", "Mobile", "Mobile Verified", "Email", "Address",
            "ID Number", "Blood Group", "Intake Channel", "Admission Type", "Profile Status",
            "Department", "Ward", "Chief Complaint", "Medico-Legal", "FIR Number",
            "Consult Fee", "Payment Status", "Receipt Number", "Admitted At", "Discharged At", "Created At",
        ],
        [
            [
                str(p.id), p.full_name, _fmt(p.date_of_birth), p.gender.value, p.mobile, p.mobile_verified,
                p.email, p.address, p.id_number, p.blood_group, p.intake_channel.value, p.admission_type.value,
                p.profile_status.value, dept_name_by_id.get(p.department_id), p.ward, p.chief_complaint,
                p.medico_legal, p.fir_number, float(p.consult_fee) if p.consult_fee is not None else None,
                p.payment_status.value, p.receipt_number, _fmt(p.admitted_at), _fmt(p.discharged_at),
                _fmt(p.created_at),
            ]
            for p in patients
        ],
    )

    # ---- Staff, split by role (doctors / nurses / other) ---------------------
    users = list(db.execute(select(User).where(User.role != Role.patient).order_by(User.role, User.full_name)).scalars())
    doctor_by_user_id = {d.user_id: d for d in db.execute(select(Doctor)).scalars()}

    doctor_rows, nurse_rows, other_rows = [], [], []
    for u in users:
        if u.role == Role.doctor:
            doc = doctor_by_user_id.get(u.id)
            doctor_rows.append(
                [
                    str(u.id), u.full_name, u.email, u.phone, u.is_active,
                    dept_name_by_id.get(doc.department_id) if doc else None,
                    doc.specialty if doc else None, _fmt(u.created_at),
                ]
            )
        elif u.role == Role.head_nurse:
            nurse_rows.append([str(u.id), u.full_name, u.email, u.phone, u.is_active, u.ward, _fmt(u.created_at)])
        else:
            other_rows.append([str(u.id), u.full_name, u.email, u.phone, u.role.value, u.is_active, _fmt(u.created_at)])

    _write_sheet(wb, "Doctors", ["ID", "Name", "Email", "Phone", "Active", "Department", "Specialty", "Created At"], doctor_rows)
    _write_sheet(wb, "Nurses", ["ID", "Name", "Email", "Phone", "Active", "Ward", "Created At"], nurse_rows)
    _write_sheet(wb, "Other Staff", ["ID", "Name", "Email", "Phone", "Role", "Active", "Created At"], other_rows)

    # ---- Doctor duty roster ---------------------------------------------------
    roster_rows = []
    doctor_name_by_id = {d.id: u.full_name for d, u in db.execute(select(Doctor, User).join(User, User.id == Doctor.user_id)).all()}
    for r in db.execute(select(DoctorRoster).order_by(DoctorRoster.doctor_id, DoctorRoster.day_of_week)).scalars():
        roster_rows.append([doctor_name_by_id.get(r.doctor_id), DAY_NAMES[r.day_of_week], r.is_on_duty, r.max_patients])
    _write_sheet(wb, "Doctor Roster", ["Doctor", "Day", "On Duty", "Max Patients"], roster_rows)

    # ---- Departments ------------------------------------------------------
    _write_sheet(
        wb, "Departments", ["ID", "Name", "Code"],
        [[str(d.id), d.name, d.code] for d in db.execute(select(Department).order_by(Department.name)).scalars()],
    )

    # ---- Patient bills (consult fee, per patient) -----------------------------
    _write_sheet(
        wb,
        "Consult Billing",
        ["Patient", "Mobile", "Consult Fee", "Payment Status", "Receipt Number", "Admitted At"],
        [
            [p.full_name, p.mobile, float(p.consult_fee) if p.consult_fee is not None else None,
             p.payment_status.value, p.receipt_number, _fmt(p.admitted_at)]
            for p in patients
            if p.consult_fee is not None
        ],
    )

    # ---- Pharmacy: stock, formulary, dispense/billing history -----------------
    _write_sheet(
        wb,
        "Pharmacy Stock",
        ["ID", "Medicine", "Batch", "Expiry", "Quantity", "Min Threshold", "Unit Price", "Expired Flagged"],
        [
            [str(s.id), s.medicine_name, s.batch_number, _fmt(s.expiry_date), s.quantity, s.min_threshold,
             float(s.unit_price), s.is_expired_flagged]
            for s in db.execute(select(StockItem).order_by(StockItem.medicine_name)).scalars()
        ],
    )
    _write_sheet(
        wb, "Medicine Formulary", ["ID", "Name", "Default Dosage", "Approved"],
        [[str(m.id), m.name, m.default_dosage, m.is_approved] for m in db.execute(select(MedicineFormulary).order_by(MedicineFormulary.name)).scalars()],
    )

    patient_name_by_id = {p.id: p.full_name for p in patients}
    billing_rows = []
    for entry, log in db.execute(
        select(BillingEntry, DispenseLog).outerjoin(DispenseLog, BillingEntry.dispense_log_id == DispenseLog.id).order_by(BillingEntry.created_at.desc())
    ).all():
        billing_rows.append(
            [
                patient_name_by_id.get(entry.patient_id, "—"), entry.description, log.quantity if log else None,
                float(entry.amount), entry.payment_status.value, entry.receipt_number, _fmt(entry.created_at), _fmt(entry.paid_at),
            ]
        )
    _write_sheet(
        wb, "Pharmacy Billing",
        ["Patient", "Medicine", "Quantity", "Amount", "Payment Status", "Receipt Number", "Dispensed At", "Paid At"],
        billing_rows,
    )

    # ---- Labs: test catalogue + all test orders ------------------------------
    _write_sheet(
        wb, "Test Catalogue", ["ID", "Name", "Department", "Category", "TAT Min (h)", "TAT Max (h)", "Price"],
        [
            [str(t.id), t.name, dept_name_by_id.get(t.department_id), t.category, float(t.tat_min_hours), float(t.tat_max_hours), float(t.price)]
            for t in db.execute(select(TestCatalogue).order_by(TestCatalogue.name)).scalars()
        ],
    )
    test_name_by_id = {t.id: t.name for t in db.execute(select(TestCatalogue)).scalars()}
    lab_rows = []
    for o in db.execute(select(TestOrder).order_by(TestOrder.ordered_at.desc())).scalars():
        lab_rows.append(
            [
                patient_name_by_id.get(o.patient_id, "—"), test_name_by_id.get(o.test_type_id, "—"), o.status.value,
                float(o.amount), o.payment_status.value, o.receipt_number, _fmt(o.ordered_at), _fmt(o.completed_at), _fmt(o.paid_at),
            ]
        )
    _write_sheet(
        wb, "Lab Billing",
        ["Patient", "Test", "Status", "Amount", "Payment Status", "Receipt Number", "Ordered At", "Completed At", "Paid At"],
        lab_rows,
    )

    # ---- Staff salaries (every period on record) -------------------------------
    user_name_by_id = {u.id: u.full_name for u in users}
    salary_rows = []
    for s in db.execute(select(StaffSalary).order_by(StaffSalary.period.desc(), StaffSalary.user_id)).scalars():
        salary_rows.append(
            [user_name_by_id.get(s.user_id, "—"), s.period, float(s.amount), s.status.value, _fmt(s.paid_at), s.notes]
        )
    _write_sheet(wb, "Staff Salaries", ["Staff", "Period", "Amount", "Status", "Paid At", "Notes"], salary_rows)

    # ---- Audit log (most recent 1000 actions) -----------------------------
    audit_rows = []
    for a in db.execute(select(AuditLog).order_by(AuditLog.created_at.desc()).limit(1000)).scalars():
        audit_rows.append(
            [user_name_by_id.get(a.user_id, "—") if a.user_id else "system", a.action, a.entity, a.entity_id, _fmt(a.created_at)]
        )
    _write_sheet(wb, "Audit Log (last 1000)", ["User", "Action", "Entity", "Entity ID", "At"], audit_rows)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
